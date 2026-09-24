import sys
import os

# Create sample DOCX file
from docx import Document

doc = Document()
doc.add_heading('Deep Learning & Explainable AI Research Report (2026)', 0)

doc.add_heading('1. Executive Abstract', level=1)
doc.add_paragraph(
    'This research investigates the integration of Retrieval-Augmented Generation (RAG) '
    'with Explainable AI (XAI) confidence calibration. Traditional LLMs operate as black-box '
    'engines, whereas DocuExplain AI provides mathematical confidence metrics, exact source citations, '
    'and evidence span alignment.'
)

doc.add_heading('2. Evaluation Benchmarks & Milestone Deadlines', level=1)
doc.add_paragraph(
    'The project deadline for submission is August 15, 2026. The Viva Voce defense examination '
    'is scheduled for August 25, 2026. The evaluation criteria assigns 40% weightage to Explainable AI '
    'transparency, 30% to retrieval performance, 20% to UX design, and 10% to code security.'
)

doc.add_heading('3. Multi-Format Ingestion Strategy', level=1)
doc.add_paragraph(
    'The system natively ingests PDF, Microsoft Word DOCX, PowerPoint PPTX, and plain text files. '
    'Each document undergoes a 7-stage processing lifecycle: Uploading -> Text Extraction -> Cleaning -> '
    'Semantic Chunking -> Embeddings Generation -> Vector Indexing -> Ready.'
)

docx_path = "C:/Users/devag/.gemini/antigravity/scratch/docuexplain-ai/backend/sample_docs/AI_Research_Report.docx"
doc.save(docx_path)
print("Created sample DOCX at:", docx_path)
