from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.625)

# Slide 1: Title Slide
slide1 = prs.slides.add_slide(prs.slide_layouts[0])
slide1.shapes.title.text = "DocuExplain AI: Final Project Defense (2026)"
slide1.placeholders[1].text = "AI Document Intelligence Platform with Explainable AI (XAI)\nDepartment of Computer Science & Artificial Intelligence"

# Slide 2: Problem Statement & Objectives
slide2 = prs.slides.add_slide(prs.slide_layouts[1])
slide2.shapes.title.text = "1. Problem Statement & Research Objective"
tf2 = slide2.placeholders[1].text_frame
tf2.text = "Traditional LLMs operate as uninterpretable black boxes that frequently hallucinate."
p2_1 = tf2.add_paragraph()
p2_1.text = "• Objective: Build an explainable document intelligence platform combining RAG with mathematical confidence calibration."
p2_2 = tf2.add_paragraph()
p2_2.text = "• Core Concept: AI Answer + Evidence Passage + Source Attribution + Confidence Score + Explainability Audit."

# Slide 3: Submission Deadlines & Milestones
slide3 = prs.slides.add_slide(prs.slide_layouts[1])
slide3.shapes.title.text = "2. Project Deadlines & Academic Milestones"
tf3 = slide3.placeholders[1].text_frame
tf3.text = "Key timeline milestones for academic evaluation:"
p3_1 = tf3.add_paragraph()
p3_1.text = "• Milestone 1 (System Architecture & UI): July 10, 2026."
p3_2 = tf3.add_paragraph()
p3_2.text = "• Milestone 2 (RAG & ChromaDB Indexing): July 28, 2026."
p3_3 = tf3.add_paragraph()
p3_3.text = "• Final Project Submission & Demonstration Deadline: August 15, 2026."
p3_4 = tf3.add_paragraph()
p3_4.text = "• Viva Voce Examination & Oral Defense: August 25, 2026."

# Slide 4: Evaluation Weightage Criteria
slide4 = prs.slides.add_slide(prs.slide_layouts[1])
slide4.shapes.title.text = "3. Evaluation Criteria & Weightage Breakdown"
tf4 = slide4.placeholders[1].text_frame
tf4.text = "Academic grading criteria distribution:"
p4_1 = tf4.add_paragraph()
p4_1.text = "• Explainable AI (XAI) Transparency & Metric Fidelity: 40% weightage."
p4_2 = tf4.add_paragraph()
p4_2.text = "• System Architecture & Vector Retrieval Performance: 30% weightage."
p4_3 = tf4.add_paragraph()
p4_3.text = "• User Experience & 3-Column Interface Quality: 20% weightage."
p4_4 = tf4.add_paragraph()
p4_4.text = "• Code Engineering, Security & Multi-Tenant Isolation: 10% weightage."

# Slide 5: System Features & Capabilities
slide5 = prs.slides.add_slide(prs.slide_layouts[1])
slide5.shapes.title.text = "4. Key Engineering & Research Deliverables"
tf5 = slide5.placeholders[1].text_frame
tf5.text = "Enterprise-grade capabilities built into the platform:"
p5_1 = tf5.add_paragraph()
p5_1.text = "• Multi-Format Parsing: PDF, Word (DOCX), PowerPoint (PPTX), and TXT."
p5_2 = tf5.add_paragraph()
p5_2.text = "• 7-Stage Pipeline: Ingest -> Text Extraction -> Cleaning -> Chunking -> Embeddings -> Vector Indexing -> Ready."
p5_3 = tf5.add_paragraph()
p5_3.text = "• Interactive Document Evidence Viewer with passage highlighting."
p5_4 = tf5.add_paragraph()
p5_4.text = "• Semantic Vector Search & Multi-Mode Summarizer (6 formats)."

pptx_path = "C:/Users/devag/.gemini/antigravity/scratch/docuexplain-ai/backend/sample_docs/DocuExplain_AI_Defense_Slides.pptx"
prs.save(pptx_path)
print("Created sample PPTX at:", pptx_path)
